#!/usr/bin/env python3
"""Сборка сводного бюджета из типовых файлов ПланФакт (БДР по ЖК)."""

from __future__ import annotations

import argparse
import re
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

DATA_COLS = list(range(2, 10))  # B..I
DASH = "-"

# Соответствие строк шаблона (лист АПРИ) и нормализованных названий статей.
TEMPLATE_ROWS: list[tuple[int, str, str | None]] = [
    (1, "ПланФакт", None),
    (3, 'ЖК "Твоя Привелегия"', None),
    (4, "Выручка", "Доходы"),
    (5, "Нераспределенный доход", None),
    (6, "Выручка", None),
    (7, "Реализация услуг Клининг (ОД)", None),
    (8, "Реализация услуг Клининг (Доп.услуги)", None),
    (9, "Реализация услуг Клининг (снег)", None),
    (10, "Прямые расходы", None),
    (11, "Нераспределенный расход", None),
    (12, "Транспортные расходы", None),
    (13, "Переменные", None),
    (14, "Услуги техниги для уборки снега (ПТ)", None),
    (15, "Услуги садовника (ПТ)", None),
    (16, "Сыпучие материалы (ПТ)", None),
    (17, "Расходные материалы и инстументы (МОП)", None),
    (18, "Расходные материалы и инстументы  (ПТ)", None),
    (19, "Арена оборудования (ПТ)", None),
    (20, "Арена оборудования (МОП)", None),
    (21, "Заработная плата  (МОП)", None),
    (22, "Заработная плата (ПТ)", None),
    (23, "Налоги с ФОТ (МОП)", None),
    (24, "Налоги с ФОТ (ПТ)", None),
    (25, "Премми  (МОП)", None),
    (26, "Премми (ПТ)", None),
    (27, "Униформа (МОП)", None),
    (28, "Униформа (ПТ)", None),
    (29, "Ремонт оборудования (МОП)", None),
    (30, "Ремонт оборудования (ПТ)", None),
    (31, "Вывоз мусора", None),
    (32, "Штрафы, пени, неустойки", None),
    (33, "Расходы к удержанию", None),
    (34, "Общехозяйственные", None),
    (35, "Аутсорс (направления)", None),
    (36, "Заработная плата МО АУП", None),
    (37, "Инструменты: ГСМ", None),
    (38, "Медосмотры", None),
    (39, "Налоги с ФОТ (ИП)", None),
    (40, "Приобритение мебели, бытовой техники", None),
    (41, "Комиссия за перевод ЗП", None),
    (42, "Аренда квартиры (включая коммуналку)", None),
    (43, "Валовая прибыль", None),
    (44, "Валовая рентабельность", None),
    (45, "Косвенные расходы", None),
    (46, "Расход БЭК-ОФИС", None),
    (47, "Расходы на бэк-офис", None),
    (48, "Автомобиль", None),
    (49, "Автомобиль: ГСМ", None),
    (50, "Автомобиль: страхование", None),
    (51, "Автомобиль: ТО", None),
    (52, "Офисные расходы", None),
    (53, "Аренда офиса", None),
    (54, "Канцелярия и обслуживание", None),
    (55, "Интернет и связь", None),
    (56, "Банковские услуги", None),
    (57, "Заработная плата АУП (БО)", None),
    (58, "Заработная плата бухгалтерии (БО)", None),
    (59, "Заработная плата СБ (БО)", None),
    (60, "Командировочные расходы", None),
    (61, "Командировнчеы расходы: билеты (БО)", None),
    (62, "Командировнчеы расходы: проживание (БО)", None),
    (63, "Комиссия УК ДОМ", None),
    (64, "Налоги с ФОТ (ИП)_БО", None),
    (65, "Налоги с ФОТ (НДФЛ) БО", None),
    (66, "Налоги с ФОТ (страховые взносы) БО", None),
    (67, "Нотариальные услуги", None),
    (68, "Обслуживание техники (БО)", None),
    (69, "Программное обеспечение (БО)", None),
    (70, "ПланФакт", None),
    (71, "ТБ", None),
    (72, "Контур", None),
    (73, "Platrum", None),
    (74, "Расходы на сайт (БО)", None),
    (75, "Расходы на технику до 40т (БО)", None),
    (76, "Такси", None),
    (77, "ФД Аутосрс (управление)", None),
    (78, "Копортативные мероприятия", None),
    (79, "HR (аутсорс)", None),
    (80, "Операционная прибыль", None),
    (81, "Операционная рентабельность", None),
    (82, "Прочие доходы", None),
    (83, "Доход БЭК-ОФИС", None),
    (84, "Пополненин бэк-офис", None),
    (85, "Проценты по вкладам", None),
    (86, "Курсовая разница (+)", None),
    (87, "Прочие расходы", None),
    (88, "Курсовая разница (-)", None),
    (89, "EBITDA", None),
    (90, "Рентабельность по EBITDA", None),
    (91, "Амортизация", None),
    (92, "Амортизация ОС (ПТ)", None),
    (93, "Амортизация ОС (МОП)", None),
    (94, "Проценты по кредитам и займам", None),
    (95, "Налог на прибыль", None),
    (96, "Налоги УСН", None),
    (97, "Налоги НДС 5%", None),
    (98, "Чистая прибыль (убыток)", None),
    (99, "Рентабельность чистой прибыли", None),
]

ROW_CHILDREN: dict[int, list[int]] = {
    4: [5, 6],
    6: [7, 8, 9],
    10: [11, 12, 13, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 32, 33, 34],
    13: [14, 15, 16, 17, 18, 19, 20],
    34: [35, 36, 37, 38, 39, 40, 41, 42],
    45: [46, 47, 48, 52, 56, 57, 58, 59, 60, 63, 64, 65, 66, 67, 68, 69, 70, 71, 72, 73, 74, 75, 76, 77, 78, 79],
    48: [49, 50, 51],
    52: [53, 54, 55],
    60: [61, 62],
    82: [83, 84, 85, 86],
    87: [88],
    91: [92, 93],
    95: [96, 97],
}

CALC_ROWS = {43, 44, 80, 81, 89, 90, 98, 99}
SUM_ROWS = {4, 6, 10, 13, 34, 45, 48, 52, 60, 82, 87, 91, 95}
LEAF_ROWS = {row for row, _, _ in TEMPLATE_ROWS} - SUM_ROWS - CALC_ROWS - {1, 3}
PERCENT_ROWS = {44, 81, 90, 99}

LABEL_ALIASES = {
    "Расходные материалы и инстументы (ПТ)": "Расходные материалы и инстументы  (ПТ)",
    "Заработная плата (МОП)": "Заработная плата  (МОП)",
    "Премми (МОП)": "Премми  (МОП)",
    "Чистая прибыль/убыток": "Чистая прибыль (убыток)",
    "Рентабельность, %": "Рентабельность чистой прибыли",
}


def norm_label(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).strip())
    return text or None


def is_dash(value: Any) -> bool:
    return value in (None, "", DASH, "-")


def to_number(value: Any) -> float | None:
    if is_dash(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(" ", "").replace(",", ".")
        if text.endswith("%"):
            return None
        try:
            return float(text)
        except ValueError:
            return None
    return None


def sheet_title_from_filename(path: Path) -> str:
    name = path.stem
    mapping = {
        "Астон_Движение": "Астон.Движение",
        "Астон_Реформа": "Астон.Реформа",
        "Дом Милый Дом": "Милый дом",
        "Дом Милы": "Милый дом",
        "ДМД": "Милый дом",
        "River Park": "River Park",
    }
    for key, title in mapping.items():
        if key in name:
            return title[:31]
    if "ДМД" in name or "Милый" in name or "Милы" in name:
        return "Милый дом"
    for marker in ['ЖК _', 'ЖК "']:
        if marker in name:
            part = name.split(marker, 1)[1]
            part = part.split("_", 1)[0].split('"', 1)[0].strip()
            return part[:31]
    return path.stem[:31]


def resolve_template_path(input_dir: Path) -> Path:
    bundled = Path(__file__).parent / "template" / "bdr_template.xlsx"
    if bundled.exists():
        return bundled
    ekaterinburg = Path(__file__).parent / "input" / "Екатеринбург"
    if ekaterinburg.exists():
        apri = next(ekaterinburg.glob("*АПРИ*"), None)
        if apri:
            return apri
    apri_in_input = next(input_dir.glob("**/*АПРИ*"), None)
    if apri_in_input:
        return apri_in_input
    files = sorted(input_dir.glob("**/*.xlsx"))
    if not files:
        raise SystemExit(f"Нет .xlsx файлов в {input_dir}")
    return files[0]


def copy_cell_style(src, dst) -> None:
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def build_source_index(ws) -> dict[str, list[tuple[int, dict[int, Any]]]]:
    index: dict[str, list[tuple[int, dict[int, Any]]]] = {}
    for row in range(1, ws.max_row + 1):
        label = norm_label(ws.cell(row, 1).value)
        if not label:
            continue
        values = {col: ws.cell(row, col).value for col in DATA_COLS}
        index.setdefault(label, []).append((row, values))
        alias = LABEL_ALIASES.get(label)
        if alias:
            index.setdefault(alias, []).append((row, values))
    return index


def pick_source_values(
    index: dict[str, list[tuple[int, dict[int, Any]]]],
    label: str,
    alias: str | None,
    prefer_row: int | None = None,
) -> dict[int, Any] | None:
    candidates: list[tuple[int, dict[int, Any]]] = []
    for key in [label, alias]:
        if key and key in index:
            candidates.extend(index[key])
    if not candidates:
        return None
    if prefer_row is not None:
        for src_row, values in candidates:
            if src_row == prefer_row:
                return values
    return candidates[0][1]


def extract_row_values(ws) -> dict[int, dict[int, Any]]:
    index = build_source_index(ws)
    rows: dict[int, dict[int, Any]] = {}

    for row, label, alias in TEMPLATE_ROWS:
        if row in CALC_ROWS or row in SUM_ROWS:
            continue
        values = pick_source_values(index, label, alias)
        if values is None:
            rows[row] = {col: DASH for col in DATA_COLS}
        else:
            rows[row] = dict(values)

    # В формате «Доходы/Расходы» выручка часто только в строке «Доходы» / «Выручка».
    if all(is_dash(rows.get(7, {}).get(col)) for col in DATA_COLS):
        for label in ("Реализация услуг Клининг (ОД)", "Выручка", "Доходы"):
            alt = pick_source_values(index, label, None)
            if alt and any(not is_dash(alt.get(col)) for col in DATA_COLS):
                rows[7] = dict(alt)
                break

    return rows


def load_template_ws(template_path: Path):
    wb = openpyxl.load_workbook(template_path)
    return wb, wb.active


def cell_ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def sum_formula(children: list[int], col: int) -> str:
    refs = ",".join(cell_ref(child, col) for child in children)
    return f"=SUM({refs})"


def cross_sheet_sum_formula(sheet_names: list[str], row: int, col: int) -> str:
    refs = ",".join(f"'{name}'!{cell_ref(row, col)}" for name in sheet_names)
    return f"=SUM({refs})"


def apply_calc_formulas(ws, col: int) -> None:
    c = cell_ref
    ws.cell(43, col).value = f"={c(4, col)}-{c(10, col)}"
    ws.cell(44, col).value = f'=IF({c(4, col)}=0,"-",{c(43, col)}/{c(4, col)})'
    ws.cell(80, col).value = f"={c(43, col)}-{c(45, col)}"
    ws.cell(81, col).value = f'=IF({c(4, col)}=0,"-",{c(80, col)}/{c(4, col)})'
    ws.cell(89, col).value = f"={c(80, col)}+{c(82, col)}-{c(87, col)}"
    ws.cell(90, col).value = f'=IF({c(4, col)}=0,"-",{c(89, col)}/{c(4, col)})'
    ws.cell(98, col).value = f"={c(89, col)}-{c(91, col)}-{c(94, col)}-{c(95, col)}"
    ws.cell(99, col).value = f'=IF({c(4, col)}=0,"-",{c(98, col)}/{c(4, col)})'


def apply_sum_formulas(ws) -> None:
    for parent, children in ROW_CHILDREN.items():
        for col in DATA_COLS:
            ws.cell(parent, col).value = sum_formula(children, col)
    for col in DATA_COLS:
        apply_calc_formulas(ws, col)
    for row in PERCENT_ROWS:
        for col in DATA_COLS:
            ws.cell(row, col).number_format = "0.00%"


def write_leaf_values(ws, rows: dict[int, dict[int, Any]]) -> None:
    for row in LEAF_ROWS:
        for col in DATA_COLS:
            ws.cell(row, col).value = rows.get(row, {}).get(col, DASH)


def write_sheet(
    ws,
    template_ws,
    rows: dict[int, dict[int, Any]],
    header_label: str | None = None,
) -> None:
    for row in range(1, template_ws.max_row + 1):
        for col in range(1, template_ws.max_column + 1):
            src = template_ws.cell(row, col)
            dst = ws.cell(row, col)
            dst.value = src.value
            copy_cell_style(src, dst)

    if header_label:
        ws.cell(3, 1).value = header_label

    write_leaf_values(ws, rows)
    apply_sum_formulas(ws)


def write_consolidated_sheet(
    ws,
    template_ws,
    project_sheet_names: list[str],
    header_label: str,
) -> None:
    for row in range(1, template_ws.max_row + 1):
        for col in range(1, template_ws.max_column + 1):
            src = template_ws.cell(row, col)
            dst = ws.cell(row, col)
            dst.value = src.value
            copy_cell_style(src, dst)

    ws.cell(3, 1).value = header_label

    for row in LEAF_ROWS:
        for col in DATA_COLS:
            ws.cell(row, col).value = cross_sheet_sum_formula(project_sheet_names, row, col)

    apply_sum_formulas(ws)


def build_workbook(
    input_dir: Path,
    output_path: Path,
    consolidated_title: str,
    header_label: str | None = None,
) -> None:
    files = sorted(input_dir.glob("**/*.xlsx"))
    if not files:
        raise SystemExit(f"Нет .xlsx файлов в {input_dir}")

    if header_label is None:
        header_label = consolidated_title

    template_file = resolve_template_path(input_dir)
    template_wb, template_ws = load_template_ws(template_file)

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    project_sheet_names: list[str] = []
    used_titles: set[str] = set()

    for path in files:
        src_wb = openpyxl.load_workbook(path, data_only=True)
        rows = extract_row_values(src_wb.active)

        title = sheet_title_from_filename(path)
        base = title
        n = 2
        while title in used_titles:
            suffix = f" {n}"
            title = (base[: 31 - len(suffix)] + suffix)
            n += 1
        used_titles.add(title)
        project_sheet_names.append(title)

        ws = out_wb.create_sheet(title=title)
        write_sheet(ws, template_ws, rows, header_label=title)
        src_wb.close()

    summary = out_wb.create_sheet(title=consolidated_title[:31], index=0)
    write_consolidated_sheet(summary, template_ws, project_sheet_names, header_label)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(output_path)
    template_wb.close()
    out_wb.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Сборка сводного бюджета ЖК")
    parser.add_argument("--input", type=Path, default=Path(__file__).parent / "input" / "Екатеринбург")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).parent / "output" / "Сводный бюджет Екатеринбурга.xlsx",
    )
    parser.add_argument("--title", default="Сводный бюджет Екатеринбурга")
    parser.add_argument("--header", default=None, help="Заголовок в ячейке A3 сводного листа")
    args = parser.parse_args()
    build_workbook(args.input, args.output, args.title, header_label=args.header)
    print(f"Готово: {args.output}")


if __name__ == "__main__":
    main()
